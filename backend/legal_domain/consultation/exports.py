"""Portable in-memory Word/PDF exports from the same public report snapshot."""
from html import escape
from io import BytesIO
import re

from .reporting import report_markdown

MIME_TYPES = {'docx': 'application/vnd.openxmlformats-officedocument.wordprocessingml.document',
              'pdf': 'application/pdf', 'md': 'text/markdown; charset=utf-8'}
LINK = re.compile(r'\[([^\]]+)\]\((https://[^\s)]+)\)')


def blocks(markdown):
    for line in markdown.splitlines():
        line = re.sub(r'[\x00-\x08\x0b\x0c\x0e-\x1f]', '', line).strip()
        if not line:
            continue
        heading = re.match(r'^(#{1,3})\s+(.*)', line)
        if heading:
            yield 'h' + str(len(heading[1])), heading[2]
        elif line.startswith('- '):
            yield 'li', line[2:]
        else:
            yield 'p', line


def _add_docx_runs(paragraph, text):
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.opc.constants import RELATIONSHIP_TYPE as RT
    cursor = 0
    for match in LINK.finditer(text):
        paragraph.add_run(text[cursor:match.start()].replace('**', ''))
        hyperlink = OxmlElement('w:hyperlink')
        hyperlink.set(qn('r:id'), paragraph.part.relate_to(match[2], RT.HYPERLINK, is_external=True))
        run = OxmlElement('w:r')
        properties = OxmlElement('w:rPr')
        color = OxmlElement('w:color')
        color.set(qn('w:val'), '174E64')
        properties.append(color)
        run.append(properties)
        label = OxmlElement('w:t')
        label.text = match[1]
        run.append(label)
        hyperlink.append(run)
        paragraph._p.append(hyperlink)
        cursor = match.end()
    paragraph.add_run(text[cursor:].replace('**', ''))


def _word(markdown: str) -> bytes:
    from docx import Document
    from docx.shared import Cm, Pt, RGBColor
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    document = Document()
    # Some installed python-docx templates inherit a colored Title border.
    # Remove decorative paragraph borders rather than relying on theme defaults.
    for border in document.styles.element.xpath('.//w:pBdr'):
        border.getparent().remove(border)
    section = document.sections[0]
    section.page_width, section.page_height = Cm(21), Cm(29.7)
    section.top_margin, section.bottom_margin = Cm(2), Cm(2)
    section.left_margin = section.right_margin = Cm(2.2)
    for name, size in [('Normal', 10.5), ('Title', 22), ('Heading 1', 15), ('Heading 2', 12), ('List Bullet', 10.5)]:
        style = document.styles[name]
        style.font.name = 'Microsoft YaHei'
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor(0, 0, 0)
        rpr = style.element.get_or_add_rPr()
        fonts = rpr.find(qn('w:rFonts'))
        if fonts is None:
            fonts = OxmlElement('w:rFonts')
            rpr.append(fonts)
        fonts.set(qn('w:eastAsia'), 'Microsoft YaHei')
        style.paragraph_format.space_after = Pt(4)
        style.paragraph_format.line_spacing = 1.15
        if name in ('Title', 'Heading 1', 'Heading 2'):
            style.paragraph_format.keep_with_next = True
    for kind, text in blocks(markdown):
        style = {'h1': 'Title', 'h2': 'Heading 1', 'h3': 'Heading 2', 'li': 'List Bullet'}.get(kind)
        paragraph = document.add_paragraph(style=style)
        _add_docx_runs(paragraph, text)
    footer = section.footer.paragraphs[0]
    footer.alignment = 2
    footer.add_run('LexPilot 律策  ·  ')
    field = OxmlElement('w:fldSimple')
    field.set(qn('w:instr'), 'PAGE')
    footer._p.append(field)
    document.core_properties.title = '法律咨询行动方案'
    document.core_properties.author = 'LexPilot 律策'
    output = BytesIO()
    document.save(output)
    return output.getvalue()


def _inline_html(text):
    parts, cursor = [], 0
    for match in LINK.finditer(text):
        parts += [escape(text[cursor:match.start()].replace('**', '')),
                  f'<a href="{escape(match[2], quote=True)}">{escape(match[1])}</a>']
        cursor = match.end()
    parts.append(escape(text[cursor:].replace('**', '')))
    return ''.join(parts)


def _pdf(markdown: str) -> bytes:
    # MuPDF's bundled CJK fallback font is embedded in the PDF. No browser,
    # Office installation, remote font or external document upload is needed.
    import pymupdf
    content = list(blocks(markdown))
    css = 'body {font-family: sans-serif; font-size: 10.5pt; line-height: 1.5; color: #202b34;} h1 {font-size: 21pt; color: black;} h2 {font-size: 14pt; color: black; margin-top: 18pt;} h3 {font-size: 12pt; color: black; margin-top: 12pt;} p {margin: 0 0 7pt;} a {color: #174e64;}'
    media = pymupdf.paper_rect('a4')
    area = pymupdf.Rect(56, 48, media.width - 56, media.height - 48)
    def next_page(index, filled):
        if index >= 160:
            raise ValueError('Report exceeds the supported page count')
        return media, area, None
    # MuPDF does not implement CSS break-avoid. Detect orphan headings from
    # actual layout positions and move them together with the following text.
    page_breaks = set()
    for attempt in range(16):
        elements = []
        for index, (kind, text) in enumerate(content):
            tag = 'p' if kind == 'li' else kind
            prefix = '• ' if kind == 'li' else ''
            style = ' style="page-break-before: always"' if index in page_breaks else ''
            elements.append(f'<{tag} id="block-{index}"{style}>{prefix}{_inline_html(text)}</{tag}>')
        positions = {}
        def remember(position):
            if position.id and position.id.startswith('block-') and position.open_close & 1:
                positions.setdefault(int(position.id[6:]), position.page_num)
        story = pymupdf.Story(html='<html><body>' + ''.join(elements) + '</body></html>', user_css=css)
        document = story.write_with_links(next_page, positionfn=remember)
        additions = set()
        for index, (kind, _) in enumerate(content[:-1]):
            if kind.startswith('h') and positions.get(index, 0) < positions.get(index+1, 0):
                first = index
                while first > 0 and content[first-1][0].startswith('h'):
                    first -= 1
                additions.add(first)
        additions -= page_breaks
        if not additions or attempt == 15:
            break
        document.close()
        # An earlier move changes all later pages; reflow before deciding the
        # next break instead of leaving unnecessary sparse pages at the end.
        page_breaks.add(min(additions))
    with document:
        document.set_metadata({'title': '法律咨询行动方案', 'author': 'LexPilot 律策'})
        for i, page in enumerate(document):
            page.insert_text((56, media.height - 26), f'LexPilot  |  {i+1} / {len(document)}', fontsize=8, color=(0.4, 0.4, 0.4))
        return document.tobytes(garbage=4, deflate=True)


def export_report(state, format: str) -> bytes:
    if format not in MIME_TYPES:
        raise ValueError('Supported formats: docx, pdf, md')
    if not state.final_report:
        raise ValueError('Generate a report before downloading')
    markdown = report_markdown(state)
    if format == 'md':
        return markdown.encode('utf-8')
    return _word(markdown) if format == 'docx' else _pdf(markdown)
